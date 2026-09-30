"""Corpus ingest harness — load the thesis corpus into a running backend.

Purpose
-------
Read the 200 synthetic/real incident cases from the thesis corpus
(``data/Corpus Tesis.xlsx`` sheet 1 and its CSV twin), POST each one to the
RUNNING backend as an incident, and write the measured automated times back
into BOTH files plus a ``data/corpus_resultados.json`` sidecar.

Fidelity caveats (see README.md for the full list)
--------------------------------------------------
- The corpus is TEXT. There is no audio, so NO speech-to-text is exercised.
  For the ``Telefono`` channel the real flow is recording -> STT; here the
  transcript text is ingested directly.
- Classification is performed by the backend HYBRID classifier
  (deterministic rules + Gemini), NOT by the n8n AI Agent the real phone
  channel uses.
- The "automated time" written back is the client wall-clock of the POST;
  ``latencia_e2e_ms`` is the backend-derived ``persistido_en - ingresado_en``.

Privacy
-------
Incident descriptions are REAL internal text. This module NEVER logs, prints
or persists a description. Only identifiers, status codes and timings are
emitted. The sidecar JSON is description-free by construction.

No secrets live in this file. Read credentials from
``INGEST_OPERATOR_USERNAME`` / ``INGEST_OPERATOR_PASSWORD``.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

# ── Canonical corpus/file contract ──────────────────────────────────────────

LATENCIA_COL = "Latencia e2e (ms)"

HEADER_ID = "ID"
HEADER_DESCRIPCION = "Descripcion"
HEADER_CANAL = "Canal de Origen"
HEADER_AUTO_TIME_CANDIDATES = (
    "TIempo de Registro Automatico (Segundos)",  # casing as shipped in the corpus
    "Tiempo de Registro Automatico (Segundos)",  # tolerant variant
)

ORIGEN_EVENTO_CREACION = "creacion_incidente"
MESSAGE_ID_PREFIX = "corpus-"

# canal_origen catalog ids (migration 001): 1=correo, 2=formulario, 3=telefono.
CHANNEL_TO_CANAL_ID = {
    "correo": 1,
    "formulario web": 2,
    "telefono": 3,
}

DEFAULT_BASE_URL = "https://localhost"
DEFAULT_CSV = "data/Corpus Tesis - Hoja 1.csv"
DEFAULT_XLSX = "data/Corpus Tesis.xlsx"
DEFAULT_SIDECAR = "data/corpus_resultados.json"

LOGIN_PATH = "/api/v1/auth/login"
INCIDENTES_PATH = "/api/v1/incidentes/"

# Retry policy for transient 5xx / timeouts (e.g. Gemini 503 through the
# backend classifier). 4xx are never retried.
MAX_ATTEMPTS = 3
BACKOFF_BASE_S = 1.0


# ── Pure helpers (unit tested offline) ──────────────────────────────────────


def _norm(value: Any) -> str:
    """Normalize a cell value for header comparison (trim + lowercase)."""
    if value is None:
        return ""
    return str(value).strip().lower()


def map_canal_origen_id(channel: str | None) -> int | None:
    """Map a corpus channel label to the canonical ``canal_origen.id``.

    ``Correo`` -> 1, ``Formulario Web`` -> 2, ``Telefono`` -> 3. Unknown or
    empty labels return ``None`` so the caller can fail that case explicitly.
    Matching is case- and whitespace-insensitive.
    """
    return CHANNEL_TO_CANAL_ID.get(_norm(channel))


def build_incidente_payload(
    *,
    descripcion: str,
    canal_origen_id: int | None,
    case_id: str,
    ingresado_en: datetime | None = None,
) -> dict[str, Any]:
    """Build the ``POST /api/v1/incidentes/`` body for one corpus case.

    ``origen_message_id`` is deterministic (``corpus-<ID>``) so a re-run is
    idempotent server-side. ``ingresado_en`` is always ISO-8601 with an
    explicit timezone, as required by the backend to derive ``latencia_e2e_ms``.
    """
    instant = (ingresado_en or datetime.now(timezone.utc)).astimezone(timezone.utc)
    return {
        "descripcion": descripcion,
        "canal_origen_id": canal_origen_id,
        "origen_evento": ORIGEN_EVENTO_CREACION,
        "origen_message_id": f"{MESSAGE_ID_PREFIX}{case_id}",
        "ingresado_en": instant.isoformat(),
    }


def _format_seconds(value: float | None) -> str:
    """Format a wall-clock measurement for a CSV cell (blank when missing)."""
    if value is None:
        return ""
    text = f"{round(float(value), 3):.3f}".rstrip("0").rstrip(".")
    return text or "0"


def _find_header_row(rows: Sequence[Sequence[Any]]) -> int:
    """Return the index of the header row (first row with ID + Descripcion)."""
    for idx, row in enumerate(rows):
        normalized = {_norm(cell) for cell in row}
        if HEADER_ID.lower() in normalized and HEADER_DESCRIPCION.lower() in normalized:
            return idx
    raise ValueError("No se encontro la fila de encabezados (ID/Descripcion) en el corpus.")


def _column_index(header: Sequence[Any], name: str) -> int:
    target = _norm(name)
    for idx, cell in enumerate(header):
        if _norm(cell) == target:
            return idx
    raise ValueError(f"Columna requerida ausente en el encabezado: {name!r}")


def _auto_time_index(header: Sequence[Any]) -> int:
    for candidate in HEADER_AUTO_TIME_CANDIDATES:
        try:
            return _column_index(header, candidate)
        except ValueError:
            continue
    raise ValueError(
        "No se encontro la columna de tiempo automatico en el encabezado."
    )


def _ensure_column(rows: list[list[Any]], header_idx: int, name: str) -> int:
    """Return the column index for ``name``, appending it (all rows) if absent.

    Idempotency: an existing column with the same normalized name is reused,
    never duplicated. When created, an empty cell is appended to every row so
    the table stays rectangular.
    """
    header = rows[header_idx]
    try:
        return _column_index(header, name)
    except ValueError:
        pass
    header.append(name)
    for row in rows:
        if row is not header:
            row.append("")
    return len(header) - 1


# ── Result / case models ────────────────────────────────────────────────────


@dataclass
class CorpusCase:
    """One corpus row, parsed from CSV or XLSX."""

    case_id: str
    descripcion: str
    canal: str


@dataclass
class CorpusResult:
    """Per-case outcome persisted to the sidecar and written back to the corpus.

    Deliberately description-free. ``error`` is a short machine-ish label
    (e.g. ``http502``), never a response body.
    """

    case_id: str
    incidente_id: int | None
    numero_incidente: str | None
    sector: str | None
    revision: bool | None
    wall_s: float | None
    latencia_e2e_ms: int | None
    error: str | None


# ── Reading ─────────────────────────────────────────────────────────────────


def read_csv_cases(path: str | Path) -> list[CorpusCase]:
    """Parse corpus cases from the CSV twin."""
    with open(path, encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh))
    header_idx = _find_header_row(rows)
    header = rows[header_idx]
    id_col = _column_index(header, HEADER_ID)
    desc_col = _column_index(header, HEADER_DESCRIPCION)
    canal_col = _column_index(header, HEADER_CANAL)

    cases: list[CorpusCase] = []
    for row in rows[header_idx + 1 :]:
        if not row or not row[id_col].strip():
            continue
        cases.append(
            CorpusCase(
                case_id=row[id_col].strip(),
                descripcion=row[desc_col],
                canal=row[canal_col].strip(),
            )
        )
    return cases


def read_xlsx_cases(path: str | Path) -> list[CorpusCase]:
    """Parse corpus cases from sheet 1 of the XLSX file (openpyxl required)."""
    openpyxl = _import_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    header_idx = _find_header_row(rows)
    header = rows[header_idx]
    id_col = _column_index(header, HEADER_ID)
    desc_col = _column_index(header, HEADER_DESCRIPCION)
    canal_col = _column_index(header, HEADER_CANAL)
    cases: list[CorpusCase] = []
    for row in rows[header_idx + 1 :]:
        if not row or not str(row[id_col] or "").strip():
            continue
        cases.append(
            CorpusCase(
                case_id=str(row[id_col]).strip(),
                descripcion=str(row[desc_col] or ""),
                canal=str(row[canal_col] or "").strip(),
            )
        )
    return cases


def read_cases(path: str | Path, source: str = "csv") -> list[CorpusCase]:
    """Dispatch to the CSV or XLSX reader."""
    if source == "xlsx":
        return read_xlsx_cases(path)
    return read_csv_cases(path)


# ── Writing (idempotent, structure-preserving) ──────────────────────────────


def write_csv_results(
    path: str | Path, results: dict[str, CorpusResult]
) -> None:
    """Fill the automated-time column and (add/update) the latencia column.

    Preserves the title row, the header row and every original column/row.
    Re-running updates the same two columns and never duplicates them.
    """
    with open(path, encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh))
    header_idx = _find_header_row(rows)
    header = rows[header_idx]
    id_col = _column_index(header, HEADER_ID)
    auto_col = _auto_time_index(header)
    lat_col = _ensure_column(rows, header_idx, LATENCIA_COL)

    width = len(header)
    for row in rows:
        if len(row) < width:
            row.extend([""] * (width - len(row)))

    for row in rows[header_idx + 1 :]:
        case_id = row[id_col].strip()
        result = results.get(case_id)
        if result is None:
            continue
        row[auto_col] = _format_seconds(result.wall_s)
        row[lat_col] = "" if result.latencia_e2e_ms is None else str(result.latencia_e2e_ms)

    with open(path, "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerows(rows)


def write_xlsx_results(
    path: str | Path, results: dict[str, CorpusResult]
) -> None:
    """Update the XLSX in place, preserving all other columns and rows.

    Reuses an existing ``Latencia e2e (ms)`` column if present; otherwise adds
    it at the end. Idempotent on re-run.
    """
    openpyxl = _import_openpyxl()
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]

    header_row = None
    header_values: list[Any] = []
    for row_idx in range(1, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
        normalized = {_norm(v) for v in values}
        if HEADER_ID.lower() in normalized and HEADER_DESCRIPCION.lower() in normalized:
            header_row = row_idx
            header_values = values
            break
    if header_row is None:
        raise ValueError("No se encontro la fila de encabezados en el XLSX.")

    id_col = _column_index(header_values, HEADER_ID) + 1
    auto_col = _auto_time_index(header_values) + 1

    lat_col: int | None = None
    for idx, cell in enumerate(header_values, start=1):
        if _norm(cell) == _norm(LATENCIA_COL):
            lat_col = idx
            break
    if lat_col is None:
        lat_col = len(header_values) + 1
        ws.cell(row=header_row, column=lat_col, value=LATENCIA_COL)

    for row_idx in range(header_row + 1, ws.max_row + 1):
        case_id = ws.cell(row=row_idx, column=id_col).value
        if case_id is None or not str(case_id).strip():
            continue
        result = results.get(str(case_id).strip())
        if result is None:
            continue
        ws.cell(
            row=row_idx,
            column=auto_col,
            value=None if result.wall_s is None else round(float(result.wall_s), 3),
        )
        ws.cell(row=row_idx, column=lat_col, value=result.latencia_e2e_ms)

    wb.save(path)


def write_sidecar_json(path: str | Path, results: Iterable[CorpusResult]) -> None:
    """Write the description-free traceability sidecar."""
    payload = [asdict(r) for r in results]
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with open(target, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def _import_openpyxl():
    """Import openpyxl lazily so CSV-only paths/tests work without it."""
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "openpyxl es requerido para leer/escribir XLSX. "
            "Instalar con: pip install -r scripts/corpus_ingest/requirements.txt"
        ) from exc
    return openpyxl


# ── HTTP layer ──────────────────────────────────────────────────────────────


def _request_with_retry(
    session: Any,
    method: str,
    url: str,
    *,
    max_attempts: int = MAX_ATTEMPTS,
    backoff_base: float = BACKOFF_BASE_S,
    **kwargs: Any,
) -> Any:
    """Send a request, retrying only 5xx responses and network timeouts.

    4xx responses are returned immediately (they are deterministic and must not
    be retried). Exhausted retries re-raise the last network exception.
    """
    import requests  # local import: keeps module import side-effect free

    last_exc: Exception | None = None
    for attempt in range(1, max_attempts + 1):
        try:
            response = session.request(method, url, **kwargs)
            if response.status_code >= 500 and attempt < max_attempts:
                time.sleep(backoff_base * (2 ** (attempt - 1)))
                continue
            return response
        except (requests.Timeout, requests.ConnectionError) as exc:
            last_exc = exc
            if attempt < max_attempts:
                time.sleep(backoff_base * (2 ** (attempt - 1)))
                continue
            raise
    if last_exc is not None:  # pragma: no cover - defensive
        raise last_exc
    raise RuntimeError("request_with_retry exhausted without response")  # pragma: no cover


def login(session: Any, base_url: str, username: str, password: str, verify: Any) -> str:
    """Authenticate and return the bearer token. Never logs credentials."""
    response = _request_with_retry(
        session,
        "POST",
        f"{base_url}{LOGIN_PATH}",
        json={"username": username, "password": password},
        verify=verify,
        timeout=30,
    )
    if response.status_code != 200:
        raise RuntimeError(f"Login fallido: HTTP {response.status_code}")
    return response.json()["access_token"]


def post_incidente(
    session: Any, base_url: str, token: str, payload: dict[str, Any], verify: Any, timeout: float
) -> tuple[Any, float]:
    """POST one incident, returning (response, wall-clock seconds).

    The wall clock is measured around the POST->response round trip: this is
    the harness "automated registration time".
    """
    started = time.perf_counter()
    response = _request_with_retry(
        session,
        "POST",
        f"{base_url}{INCIDENTES_PATH}",
        headers={"Authorization": f"Bearer {token}"},
        json=payload,
        verify=verify,
        timeout=timeout,
    )
    wall_s = time.perf_counter() - started
    return response, wall_s


def _result_from_response(case_id: str, response: Any, wall_s: float) -> CorpusResult:
    """Map a successful HTTP response to a CorpusResult (no body logging)."""
    data = response.json()
    sector_obj = data.get("sector") or {}
    return CorpusResult(
        case_id=case_id,
        incidente_id=data.get("id"),
        numero_incidente=data.get("numero_incidente"),
        sector=sector_obj.get("nombre"),
        revision=data.get("requiere_revision_humana"),
        wall_s=round(wall_s, 3),
        latencia_e2e_ms=data.get("latencia_e2e_ms"),
        error=None,
    )


def _error_label(response: Any) -> str:
    """Short, body-free error label for a non-2xx response."""
    code = ""
    try:
        body = response.json()
        code = (body.get("error") or {}).get("code") or ""
    except Exception:  # noqa: BLE001 - body is untrusted
        code = ""
    return f"http{response.status_code}" + (f"/{code}" if code else "")


# ── CLI / orchestration ─────────────────────────────────────────────────────


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Carga el corpus de la tesis en el backend en ejecucion y escribe "
            "los tiempos automaticos de vuelta al XLSX y al CSV. Nunca imprime "
            "descripciones."
        )
    )
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--insecure", action="store_true", help="Aceptar certificado autofirmado.")
    parser.add_argument("--csv", default=DEFAULT_CSV)
    parser.add_argument("--xlsx", default=DEFAULT_XLSX)
    parser.add_argument("--sidecar", default=DEFAULT_SIDECAR)
    parser.add_argument("--source", choices=("csv", "xlsx"), default="csv")
    parser.add_argument("--only-channel", default=None)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--sleep", type=float, default=0.5, help="Segundos entre casos.")
    parser.add_argument("--timeout", type=float, default=30.0)
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args(argv)


def _select_cases(
    cases: Sequence[CorpusCase], only_channel: str | None, limit: int | None
) -> list[CorpusCase]:
    selected = list(cases)
    if only_channel is not None:
        target = _norm(only_channel)
        selected = [c for c in selected if _norm(c.canal) == target]
    if limit is not None:
        selected = selected[: max(limit, 0)]
    return selected


def _channel_counts(cases: Sequence[CorpusCase]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for case in cases:
        counts[case.canal] = counts.get(case.canal, 0) + 1
    return counts


def run(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    source_path = args.xlsx if args.source == "xlsx" else args.csv
    cases = read_cases(source_path, args.source)
    selected = _select_cases(cases, args.only_channel, args.limit)

    counts = _channel_counts(cases)
    print(f"Casos en el corpus: {len(cases)} | seleccionados: {len(selected)}")
    for channel, count in sorted(counts.items()):
        print(f"  canal {channel!r}: {count}")

    if args.dry_run:
        mapped = sum(1 for c in selected if map_canal_origen_id(c.canal) is not None)
        unmapped = len(selected) - mapped
        print(f"DRY-RUN: mapeados={mapped} sin_mapear={unmapped}. No se ejecuto red ni escritura.")
        return 0

    username = os.environ.get("INGEST_OPERATOR_USERNAME")
    password = os.environ.get("INGEST_OPERATOR_PASSWORD")
    if not username or not password:
        print(
            "ERROR: faltan INGEST_OPERATOR_USERNAME / INGEST_OPERATOR_PASSWORD.\n"
            "Exportar credenciales de un operador antes de correr el harness. "
            "Ver README.md (no hay secretos hardcodeados).",
            file=sys.stderr,
        )
        return 2

    import requests  # local import

    base_url = args.base_url.rstrip("/")
    session = requests.Session()
    verify: Any = False if args.insecure else True
    if args.insecure:
        try:
            import urllib3  # noqa: PLC0415

            urllib3.disable_warnings()
        except Exception:  # noqa: BLE001 - warnings suppression is best-effort
            pass

    token = login(session, base_url, username, password, verify)
    print("Login OK. Iniciando ingesta secuencial...")

    results: list[CorpusResult] = []
    for index, case in enumerate(selected, start=1):
        if index > 1 and args.sleep > 0:
            time.sleep(args.sleep)
        canal_id = map_canal_origen_id(case.canal)
        if canal_id is None:
            results.append(
                CorpusResult(
                    case_id=case.case_id,
                    incidente_id=None,
                    numero_incidente=None,
                    sector=None,
                    revision=None,
                    wall_s=None,
                    latencia_e2e_ms=None,
                    error="canal_desconocido",
                )
            )
            print(f"[{index}/{len(selected)}] {case.case_id}: canal desconocido, omitido")
            continue

        payload = build_incidente_payload(
            descripcion=case.descripcion, canal_origen_id=canal_id, case_id=case.case_id
        )
        try:
            response, wall_s = post_incidente(
                session, base_url, token, payload, verify, args.timeout
            )
            if response.status_code >= 400:
                results.append(
                    CorpusResult(
                        case_id=case.case_id,
                        incidente_id=None,
                        numero_incidente=None,
                        sector=None,
                        revision=None,
                        wall_s=round(wall_s, 3),
                        latencia_e2e_ms=None,
                        error=_error_label(response),
                    )
                )
                print(f"[{index}/{len(selected)}] {case.case_id}: error {_error_label(response)}")
            else:
                result = _result_from_response(case.case_id, response, wall_s)
                results.append(result)
                print(
                    f"[{index}/{len(selected)}] {case.case_id}: "
                    f"incidente {result.incidente_id} sector={result.sector!r} "
                    f"wall={result.wall_s}s latencia={result.latencia_e2e_ms}ms"
                )
        except Exception as exc:  # noqa: BLE001 - never leak body/description
            results.append(
                CorpusResult(
                    case_id=case.case_id,
                    incidente_id=None,
                    numero_incidente=None,
                    sector=None,
                    revision=None,
                    wall_s=None,
                    latencia_e2e_ms=None,
                    error=type(exc).__name__,
                )
            )
            print(f"[{index}/{len(selected)}] {case.case_id}: excepcion {type(exc).__name__}")

    write_sidecar_json(args.sidecar, results)
    by_id = {r.case_id: r for r in results}
    write_csv_results(args.csv, by_id)
    write_xlsx_results(args.xlsx, by_id)

    ok = sum(1 for r in results if r.error is None)
    print(f"Listo: {ok}/{len(results)} casos registrados. Sidecar: {args.sidecar}")
    return 0 if ok == len(results) else 1


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(run())
